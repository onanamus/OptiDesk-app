"""Import Winpharma (CSV / Excel) : lecture, détection des colonnes, application au catalogue.
Le protocole en deux temps (aperçu puis commit) est porté par routers/imports.py."""
import calendar
import csv
import io
import json
import re
import time
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from core.config import settings
from core.text import normalize
from models import Category, Product, utcnow

FIELDS = ("name", "code", "stock", "price", "dci", "category", "expiry")
REQUIRED = ("name", "code", "stock", "price")
FIELD_LABEL = {"name": "Désignation", "code": "Code", "stock": "Qté", "price": "Prix",
               "dci": "DCI", "category": "Catégorie", "expiry": "Péremption"}
DEFAULT_CATEGORY = "Non classé"
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000

# Synonymes d'en-têtes (sans accents, sans ponctuation), par ordre de priorité.
SYNONYMS = {
    "name": ["designation", "libelle", "produit", "nom", "intitule"],
    "code": ["codebarres", "codebarre", "codeproduit", "code", "cip", "ean", "reference"],
    "stock": ["qtestock", "quantitestock", "quantite", "qte", "stock"],
    "price": ["prixpublic", "prixdevente", "pvttc", "pvp", "prixttc", "prix"],
    "dci": ["dci", "substanceactive", "denominationcommune"],
    "category": ["famille", "rayon", "categorie", "classe"],
    "expiry": ["datedeperemption", "peremption", "dlc", "dluo", "expiration"],
}
AVOID = {"price": ("achat", "revient", "ht", "remise", "marge"),
         "stock": ("command", "min", "max", "alerte"),
         "code": ("fournisseur",)}


class ImportFileError(Exception):
    """Fichier illisible ou hors format — le message est destiné au pharmacien."""


@dataclass
class Table:
    headers: list[str]
    rows: list[list]


# --- Lecture -----------------------------------------------------------------------------------------

def _clean(cell) -> object:
    return cell.strip() if isinstance(cell, str) else cell


def _unique_headers(raw: list) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(raw):
        h = str(h).strip() if h not in (None, "") else f"Colonne {i + 1}"
        seen[h] = seen.get(h, 0) + 1
        out.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    return out


def read_table(content: bytes, filename: str) -> Table:
    ext = Path(filename).suffix.lower()
    if len(content) > MAX_BYTES:
        raise ImportFileError("Ce fichier est trop volumineux (20 Mo maximum).")
    if ext == ".xls":
        raise ImportFileError("Le format .xls (ancien Excel) n'est pas lu. "
                              "Enregistrez le fichier au format .xlsx ou .csv depuis Winpharma ou Excel.")
    if ext in (".csv", ".txt"):
        raw_rows = _read_csv(content)
    elif ext == ".xlsx":
        raw_rows = _read_xlsx(content)
    else:
        raise ImportFileError("Ce fichier n'est pas reconnu. Vérifiez qu'il vient bien de Winpharma (CSV ou Excel).")

    raw_rows = [r for r in raw_rows if any(c not in (None, "") for c in r)]
    if len(raw_rows) < 2:
        raise ImportFileError("Ce fichier ne contient aucune ligne de produit. "
                              "Vérifiez qu'il vient bien de Winpharma (CSV ou Excel).")
    if len(raw_rows) > MAX_ROWS:
        raise ImportFileError(f"Ce fichier dépasse {MAX_ROWS:,} lignes.".replace(",", " "))
    headers = _unique_headers(raw_rows[0])
    width = len(headers)
    rows = [(list(r) + [None] * width)[:width] for r in raw_rows[1:]]
    return Table(headers, rows)


def _read_csv(content: bytes) -> list[list]:
    for enc in ("utf-8-sig", "cp1252"):        # les exports Windows sont souvent en ANSI (cp1252)
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = content.decode("latin-1")
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        delimiter = ";"
    return [[_clean(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(content: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
    except Exception as exc:                        # fichier corrompu, mauvais format déguisé…
        raise ImportFileError("Ce fichier Excel est illisible. Vérifiez qu'il vient bien de Winpharma.") from exc


# --- Détection des colonnes ------------------------------------------------------------------------------

def _key(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize(header))


def detect_mapping(headers: list[str]) -> dict[str, str | None]:
    keys = {h: _key(h) for h in headers}
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    for field in FIELDS:
        found = None
        for exact in (True, False):
            for syn in SYNONYMS[field]:
                for h, k in keys.items():
                    if h in used or any(bad in k for bad in AVOID.get(field, ())):
                        continue
                    if (k == syn) if exact else (syn in k):
                        found = h
                        break
                if found:
                    break
            if found:
                break
        mapping[field] = found
        if found:
            used.add(found)
    return mapping


def missing_required(mapping: dict[str, str | None]) -> list[str]:
    return [FIELD_LABEL[f] for f in REQUIRED if not mapping.get(f)]


# --- Interprétation des cellules ------------------------------------------------------------------------------

def parse_number(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"(?i)fcfa|cfa|f\b", "", str(v)).replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):       # 1.200 = mille deux cents
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    m = re.fullmatch(r"(\d{1,2})/(\d{4})", s)         # MM/AAAA = fin de mois (usage en officine)
    if m and 1 <= int(m[1]) <= 12:
        return date(int(m[2]), int(m[1]), calendar.monthrange(int(m[2]), int(m[1]))[1])
    return None


def _code(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() if v not in (None, "") else ""


def parse_rows(table: Table, mapping: dict[str, str | None]) -> tuple[list[dict], list[str], int]:
    """(lignes valides dédoublonnées par code, lignes ignorées avec motif, nb de doublons de code)."""
    idx = {f: table.headers.index(h) for f, h in mapping.items() if h}
    valid: dict[str, dict] = {}
    skipped: list[str] = []
    duplicates = 0
    for n, row in enumerate(table.rows, start=2):        # ligne 1 = en-têtes
        name = str(row[idx["name"]] or "").strip()
        code = _code(row[idx["code"]])
        price = parse_number(row[idx["price"]])
        if not name or not code:
            skipped.append(f"ligne {n} : désignation ou code manquant")
            continue
        if price is None or price < 0:
            skipped.append(f"ligne {n} ({name}) : prix illisible")
            continue
        qty = parse_number(row[idx["stock"]])
        data = {"row": n, "name": name[:150], "code": code[:40], "price": int(round(price)),
                "stock": max(0, int(qty)) if qty is not None else 0}
        if "dci" in idx:
            data["dci"] = str(row[idx["dci"]] or "").strip()[:150]
        if "category" in idx:
            data["category"] = str(row[idx["category"]] or "").strip()[:60] or None
        if "expiry" in idx:
            data["expiry"] = parse_date(row[idx["expiry"]])
        if code in valid:
            duplicates += 1
        valid[code] = data
    return list(valid.values()), skipped, duplicates


# --- Fichiers en attente (entre l'aperçu et le commit) ----------------------------------------------------------------

def _pending_dir() -> Path:
    d = settings.data_dir / "imports"
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_pending(content: bytes, filename: str) -> str:
    d = _pending_dir()
    for old in d.glob("*"):                                  # purge des aperçus abandonnés (> 24 h)
        if time.time() - old.stat().st_mtime > 86400:
            old.unlink(missing_ok=True)
    import_id = uuid.uuid4().hex
    ext = Path(filename).suffix.lower()
    (d / f"{import_id}{ext}").write_bytes(content)
    (d / f"{import_id}.json").write_text(json.dumps({"filename": Path(filename).name[:120], "ext": ext}))
    return import_id


def load_pending(import_id: str) -> tuple[bytes, str]:
    if not re.fullmatch(r"[0-9a-f]{32}", import_id):
        raise ImportFileError("Cet aperçu d'import est introuvable. Recommencez en rechargeant le fichier.")
    meta_path = _pending_dir() / f"{import_id}.json"
    if not meta_path.exists():
        raise ImportFileError("Cet aperçu d'import a expiré. Recommencez en rechargeant le fichier.")
    meta = json.loads(meta_path.read_text())
    return (_pending_dir() / f"{import_id}{meta['ext']}").read_bytes(), meta["filename"]


def drop_pending(import_id: str) -> None:
    for f in _pending_dir().glob(f"{import_id}.*"):
        f.unlink(missing_ok=True)


# --- Application au catalogue ------------------------------------------------------------------------------------------

def _category_for(db: Session, cache: dict[str, Category], name: str | None) -> Category:
    name = name or DEFAULT_CATEGORY
    key = normalize(name)
    if key not in cache:
        cache[key] = Category(name=name)
        db.add(cache[key])
        db.flush()
    return cache[key]


def apply_import(db: Session, rows: list[dict]) -> tuple[int, int, list[str]]:
    """Upsert par code. Un produit saisi à la main n'est jamais écrasé en silence :
    il est mis à jour ET nommé dans le retour (donc dans le journal d'audit)."""
    now = utcnow()
    existing = {p.code: p for p in db.scalars(select(Product))}
    cats = {normalize(c.name): c for c in db.scalars(select(Category))}
    created = updated = 0
    manual_updated: list[str] = []
    for r in rows:
        p = existing.get(r["code"])
        if p is None:
            p = Product(code=r["code"], source="import", category=_category_for(db, cats, r.get("category")))
            db.add(p)
            created += 1
        else:
            updated += 1
            if p.source == "manuel":
                manual_updated.append(p.name if p.name == r["name"] else f"{p.name} → {r['name']}")
                p.source = "import"
            if r.get("category"):
                p.category = _category_for(db, cats, r["category"])
        p.name, p.price, p.stock, p.updated_at = r["name"], r["price"], r["stock"], now
        if "dci" in r:
            p.dci = r["dci"]
        if "expiry" in r:
            p.expiry = r["expiry"]
    return created, updated, manual_updated
